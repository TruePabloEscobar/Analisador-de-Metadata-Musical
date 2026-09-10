import hashlib
import threading
import pytest
from mutagen.id3 import ID3, RVA2
from scanner import _text, scan_one, scan_library
from updater import version_tuple


def test_rva2_unknown_channel_reproduces_dependency_bug_and_scans(tmp_path):
    tag = RVA2(desc='track', channel=16, gain=1, peak=.5)
    with pytest.raises(IndexError):
        str(tag)
    assert 'channel=16' in _text(tag)
    path = tmp_path / 'channel16.mp3'
    tags = ID3(); tags.add(tag); tags.save(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    result = scan_one(path)
    assert result.error is None
    assert any('channel=16' in t.value for t in result.raw_tags)
    assert hashlib.sha256(path.read_bytes()).hexdigest() == before


def test_unexpected_reader_error_is_recorded_and_scan_continues(tmp_path, monkeypatch):
    import scanner
    (tmp_path / 'a.mp3').write_bytes(b'x')
    (tmp_path / 'b.mp3').write_bytes(b'x')
    def fail(path):
        raise ValueError('broken metadata')
    monkeypatch.setattr(scanner, 'scan_one', fail)
    files, results, errors, cancelled, stats = scan_library(tmp_path, False, threading.Event())
    assert len(results) == len(errors) == stats.processed == 2
    assert 'Traceback' in errors[0][2]
    assert not cancelled


@pytest.mark.parametrize('version, expected', [('v1.1.0', (1,1,0)), ('1.10.0', (1,10,0))])
def test_version_comparison(version, expected):
    assert version_tuple(version) == expected


def test_prerelease_version_rejected():
    with pytest.raises(ValueError):
        version_tuple('1.2.0-beta')


def test_export_removed_audio_and_invalid_xml(tmp_path):
    from app import App
    from scanner import ScanResult
    from openpyxl import load_workbook
    path = tmp_path / 'removed.mp3'
    gui = object.__new__(App); gui.after = lambda *a: None
    result = ScanResult({'Filename':path.name, 'File Size':123, 'Title':'a\ufffeb\ud800c'})
    report = gui._export(tmp_path, ['Filename', 'Title'], [path], [(path, result)], [])
    with report.open('rb') as stream:
        book = load_workbook(stream)
        assert book['Tracks'].max_row == 2
        assert book['Tracks']['C2'].value == 'a\\xfffeb\\xd800c'
        book.close()


def test_updater_checks_size_hash_and_origin(tmp_path, monkeypatch):
    import io
    import updater
    payload = b'MZ' + b'test release'
    monkeypatch.setattr(updater.urllib.request, 'urlopen', lambda *a, **kw: io.BytesIO(payload))
    asset = {'browser_download_url': updater.PROJECT_URL + '/releases/download/v1.2.0/Track.Metadata.Scanner.exe', 'size': len(payload), 'digest': 'sha256:' + hashlib.sha256(payload).hexdigest()}
    staged = updater.download_update({'scanner_asset': asset}, tmp_path / 'scanner.exe')
    assert staged.read_bytes() == payload
    asset['digest'] = 'sha256:' + '0'*64
    with pytest.raises(ValueError, match='integridade'):
        updater.download_update({'scanner_asset': asset}, tmp_path / 'scanner.exe')
    asset['browser_download_url'] = 'https://other.example/update.exe'
    with pytest.raises(ValueError, match='Origem'):
        updater.download_update({'scanner_asset': asset}, tmp_path / 'scanner.exe')
