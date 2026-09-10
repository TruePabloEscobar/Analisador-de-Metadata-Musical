import json
import os
import sys
import threading
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from scanner import GROUPS, scan_library
from app import App


# Defina TRACK_SCANNER_TEST_LIBRARY somente no computador onde o teste será executado.
source_value = os.environ.get("TRACK_SCANNER_TEST_LIBRARY")

if not source_value:
    print(
        "TRACK_SCANNER_TEST_LIBRARY não definida. "
        "Teste com biblioteca real ignorado.",
        flush=True,
    )
    raise SystemExit(0)

source = Path(source_value).expanduser().resolve()

if not source.exists():
    print(
        f"Biblioteca de teste não encontrada: {source}",
        flush=True,
    )
    raise SystemExit(1)

if not source.is_dir():
    print(
        f"TRACK_SCANNER_TEST_LIBRARY não aponta para uma pasta: {source}",
        flush=True,
    )
    raise SystemExit(1)


output = Path("validation/real-1.1.0").resolve()
output.mkdir(parents=True, exist_ok=True)


# Snapshot simples para confirmar que o scanner não alterou tamanho nem timestamp.
before = {
    p: (p.stat().st_size, p.stat().st_mtime_ns)
    for p in source.iterdir()
    if p.is_file()
}


def progress(done, total, path, errors):
    if done % 1000 == 0 or done == total:
        print(done, total, errors, flush=True)


files, results, errors, _, stats = scan_library(
    source,
    False,
    threading.Event(),
    progress,
)

print("Exporting", len(results), flush=True)


app = object.__new__(App)
app.after = lambda *args: None

selected_fields = [
    field
    for group in GROUPS.values()
    for field in group
]

report = app._export(
    output,
    selected_fields,
    files,
    results,
    errors,
    stats,
)


book = load_workbook(report, read_only=True)

rows = 0

for row in book["Tracks"].iter_rows():
    assert len(row) == 40
    assert all(cell.data_type != "f" for cell in row)
    rows += 1


after = {
    p: (p.stat().st_size, p.stat().st_mtime_ns)
    for p in before
}


print(
    json.dumps(
        {
            "report": str(report),
            "tracks": rows - 1,
            "errors": len(errors),
            "sheets": book.sheetnames,
            "read_only_size_mtime": before == after,
        },
        ensure_ascii=False,
    ),
    flush=True,
)

book.close()
