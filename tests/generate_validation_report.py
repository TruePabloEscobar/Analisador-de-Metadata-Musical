from __future__ import annotations

import hashlib
import json
import shutil
import sys
import threading
import wave
from pathlib import Path

from mutagen.id3 import APIC, COMM, ID3, TKEY, TPUB, TXXX
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from app import App
from scanner import GROUPS, scan_library


VALIDATION = ROOT / "validation" / "integration"
LIBRARY = VALIDATION / "library"


def signature(path: Path):
    return path.stat().st_size, path.stat().st_mtime_ns, hashlib.sha256(path.read_bytes()).hexdigest()


def save_id3(path: Path, *frames):
    tags = ID3()
    for frame in frames:
        tags.add(frame)
    tags.save(path)


def make_wav(path: Path):
    with wave.open(str(path), "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(8000)
        wav.writeframes(b"\0\0" * 80)


if VALIDATION.exists():
    shutil.rmtree(VALIDATION)
LIBRARY.mkdir(parents=True)

first = LIBRARY / "tagged.mp3"
save_id3(first, APIC(encoding=3, mime="image/jpeg", type=3, desc="cover", data=b"x" * 1024), TKEY(encoding=3, text=["Fm"]), TXXX(encoding=3, desc="EnergyLevel", text=["7"]), COMM(encoding=3, lang="eng", desc="", text=["Energy 7"]))
second = LIBRARY / "conflict.mp3"
save_id3(second, TXXX(encoding=3, desc="ENERGY LEVEL", text=["8"]), TPUB(encoding=3, text=["Example Records"]), COMM(encoding=3, lang="eng", desc="", text=["Energy 7"]))
third = LIBRARY / "no-tags.wav"
make_wav(third)
fourth = LIBRARY / "corrupt.mp3"
fourth.write_bytes(b"not audio")
fifth = LIBRARY / "comment-only.mp3"
save_id3(fifth, COMM(encoding=3, lang="eng", desc="", text=["Energy 6"]))
(LIBRARY / "ignored.txt").write_text("not audio", encoding="utf-8")

before = {path.name: signature(path) for path in (first, second, third, fourth, fifth)}
files, results, errors, _, stats = scan_library(LIBRARY, False, threading.Event())
after = {path.name: signature(path) for path in (first, second, third, fourth, fifth)}

fields = [field for group in GROUPS.values() for field in group]
app = object.__new__(App)
app.after = lambda *args: None
report = app._export(VALIDATION, fields, files, results, errors, stats)
book = load_workbook(report, data_only=False)
tracks = list(book["Tracks"].iter_rows(values_only=True))
headers = list(tracks[0])
records = [dict(zip(headers, row)) for row in tracks[1:]]
normalized_binary = sum("[BINARY DATA" in str(value or "") for record in records for key, value in record.items() if key != "File ID")

print(json.dumps({
    "report": str(report),
    "found": stats.found,
    "supported": stats.supported,
    "ignored_extension": stats.ignored_extension,
    "tracks": len(records),
    "no_tags": sum(record.get("Scan Status") == "NO_TAGS" for record in records),
    "read_errors": len(errors),
    "energy": sum(bool(record.get("Energy")) for record in records),
    "energy_source_txxx_energylevel": sum(str(record.get("Energy Source", "")).casefold() == "txxx:energylevel" for record in records),
    "energy_source_comment": sum(record.get("Energy Source") == "COMMENT" for record in records),
    "energy_conflict": sum(record.get("Energy Conflict") is True for record in records),
    "normalized_binary": normalized_binary,
    "publisher_equals_comment": sum(bool(record.get("Comment")) and record.get("Publisher") == record.get("Comment") for record in records),
    "read_only": before == after,
}, ensure_ascii=False, indent=2))
