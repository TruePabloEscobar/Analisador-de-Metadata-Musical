import hashlib
import threading
import wave

from mutagen.id3 import APIC, COMM, ID3, TIT2, TKEY, TPE1, TPUB, TXXX
from mutagen.mp4 import MP4FreeForm
from openpyxl import load_workbook

from scanner import RawTag, ScanResult, _comment_energy, inventory_files, make_file_id, scan_library, scan_one


def save_id3(path, *frames):
    tags = ID3()
    for frame in frames:
        tags.add(frame)
    tags.save(path)


def artwork(size=1024):
    return APIC(encoding=3, mime="image/jpeg", type=3, desc="cover", data=b"x" * size)


def comment(text, desc="", lang="eng"):
    return COMM(encoding=3, lang=lang, desc=desc, text=[text])


def make_wav(path):
    with wave.open(str(path), "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(8000)
        wav.writeframes(b"\0\0" * 80)


def export_report(tmp_path, fields, files, results, errors, stats=None):
    from app import App
    app = object.__new__(App)
    app.after = lambda *args: None
    return app._export(tmp_path, fields, files, results, errors, stats)


def test_energy_comment_patterns_are_strict():
    assert _comment_energy(["Energy 8"]) == "8"
    assert _comment_energy(["Energy: 6"]) == "6"
    assert _comment_energy(["Energy = 7"]) == "7"
    assert _comment_energy(["Energy Level: 6"]) == "6"
    assert _comment_energy(["8 Energy"]) == "8"
    assert _comment_energy(["Recorded in 2018"]) is None


def test_file_id_changes_with_size(tmp_path):
    path = tmp_path / "a.mp3"
    path.write_bytes(b"a")
    first = make_file_id(path)
    path.write_bytes(b"ab")
    assert first != make_file_id(path)


def test_scan_is_read_only_size_mtime_and_sha256(tmp_path):
    path = tmp_path / "readonly.mp3"
    save_id3(path, TIT2(encoding=3, text=["Title"]), comment("Energy 7"))
    before = (path.stat().st_size, path.stat().st_mtime_ns, hashlib.sha256(path.read_bytes()).hexdigest())
    scan_one(path)
    after = (path.stat().st_size, path.stat().st_mtime_ns, hashlib.sha256(path.read_bytes()).hexdigest())
    assert before == after


def test_apic_tkey_and_comment_do_not_cross_contaminate(tmp_path):
    path = tmp_path / "test1.mp3"
    save_id3(path, artwork(), TKEY(encoding=3, text=["Fm"]), comment("Energy 7"))
    result = scan_one(path)
    assert result.values["Key"] == "Fm"
    assert result.values["Energy"] == "7"
    assert result.values["Comment"] == "Energy 7"
    assert all("[BINARY DATA" not in str(value) for key, value in result.values.items() if key not in {"Has Artwork", "Artwork Count"})


def test_apic_comment_energy_has_no_conflict(tmp_path):
    path = tmp_path / "test2.mp3"
    save_id3(path, artwork(), comment("Energy 8"))
    result = scan_one(path)
    assert result.values["Energy"] == "8"
    assert result.values["Energy Source"] == "COMMENT"
    assert result.values["Energy Conflict"] is False


def test_comment_number_without_energy_is_ignored(tmp_path):
    path = tmp_path / "test3.mp3"
    save_id3(path, artwork(), comment("Recorded in 2018"))
    result = scan_one(path)
    assert result.values["Energy"] == ""
    assert result.values["Energy Conflict"] is False


def test_explicit_and_comment_same_energy_has_no_conflict(tmp_path):
    path = tmp_path / "test4.mp3"
    save_id3(path, TXXX(encoding=3, desc="ENERGY", text=["8"]), comment("Energy 8"))
    result = scan_one(path)
    assert result.values["Energy"] == "8"
    assert result.values["Energy Source"] == "TXXX:ENERGY"
    assert result.values["Energy Conflict"] is False


def test_explicit_energy_wins_over_different_comment(tmp_path):
    path = tmp_path / "test5.mp3"
    save_id3(path, TXXX(encoding=3, desc="ENERGY", text=["8"]), comment("Energy 7"))
    result = scan_one(path)
    assert result.values["Energy"] == "8"
    assert result.values["Energy Raw"] == "8"
    assert result.values["Energy Source"] == "TXXX:ENERGY"
    assert result.values["Energy Conflict"] is True


def test_key_and_initial_key_use_independent_sources(tmp_path):
    path = tmp_path / "test6.mp3"
    save_id3(path, TKEY(encoding=3, text=["Am"]), TXXX(encoding=3, desc="INITIALKEY", text=["8A"]))
    result = scan_one(path)
    assert result.values["Key"] == "Am"
    assert result.values["Initial Key"] == "8A"


def test_tkey_does_not_fill_initial_key(tmp_path):
    path = tmp_path / "test7.mp3"
    save_id3(path, TKEY(encoding=3, text=["Am"]))
    result = scan_one(path)
    assert result.values["Key"] == "Am"
    assert result.values["Initial Key"] == ""


def test_comment_does_not_fill_publisher(tmp_path):
    path = tmp_path / "test8.mp3"
    save_id3(path, comment("Energy 6"))
    result = scan_one(path)
    assert result.values["Comment"] == "Energy 6"
    assert result.values["Publisher"] == ""


def test_tpub_is_publisher_and_comm_stays_comment(tmp_path):
    path = tmp_path / "test9.mp3"
    save_id3(path, TPUB(encoding=3, text=["Example Records"]), comment("Energy 6"))
    result = scan_one(path)
    assert result.values["Publisher"] == "Example Records"
    assert result.values["Comment"] == "Energy 6"


def test_valid_wav_without_tags_is_no_tags_not_error(tmp_path):
    path = tmp_path / "test10.wav"
    make_wav(path)
    result = scan_one(path)
    assert result.error is None
    assert result.values["Scan Status"] == "NO_TAGS"
    assert result.values["Duration"]


def test_corrupted_file_is_in_results_and_errors_and_scan_continues(tmp_path):
    bad = tmp_path / "test11.mp3"
    good = tmp_path / "good.wav"
    bad.write_bytes(b"not audio")
    make_wav(good)
    files, results, errors, stopped, stats = scan_library(tmp_path, False, threading.Event())
    assert len(files) == len(results) == 2
    assert len(errors) == stats.read_errors == 1
    assert any(path == bad and result.values["Scan Status"] == "METADATA_READ_ERROR" for path, result in results)
    assert not stopped


def test_any_binary_frame_is_raw_only(tmp_path):
    path = tmp_path / "test12.mp3"
    save_id3(path, artwork(2048), comment("Energy 7"))
    result = scan_one(path)
    textual_fields = ("Title", "Artist", "Album", "Genre", "BPM", "Key", "Initial Key", "Energy", "Publisher", "Comment")
    assert not any("[BINARY DATA" in str(result.values.get(field, "")) for field in textual_fields)
    assert any(tag.is_binary and "[BINARY DATA" in tag.value for tag in result.raw_tags)


def test_unicode_multiple_comments_and_raw_description_language(tmp_path):
    path = tmp_path / "São João.mp3"
    save_id3(path, TIT2(encoding=3, text=["Título"]), TPE1(encoding=3, text=["A", "B"]), comment("Energy 8"), comment("Comentário completo", "nota", "por"), TXXX(encoding=3, desc="INITIALKEY", text=["8A"]))
    result = scan_one(path)
    assert result.values["Title"] == "Título"
    assert "Energy 8" in result.values["Comment"] and "Comentário completo" in result.values["Comment"]
    assert any(tag.name == "TXXX:INITIALKEY" and tag.description == "INITIALKEY" for tag in result.raw_tags)
    assert any(tag.language == "por" for tag in result.raw_tags if tag.base_name == "COMM")


def test_inventory_counts_supported_and_ignored_extensions(tmp_path):
    make_wav(tmp_path / "one.wav")
    (tmp_path / "two.txt").write_text("x")
    sub = tmp_path / "sub"
    sub.mkdir()
    (sub / "three.flac").write_bytes(b"bad")
    files, stats = inventory_files(tmp_path, True)
    assert len(files) == stats.supported == 2
    assert stats.found == 3
    assert stats.ignored_extension == 1


def test_cancellation_is_reported(tmp_path):
    make_wav(tmp_path / "one.wav")
    cancel = threading.Event()
    cancel.set()
    files, results, errors, stopped, stats = scan_library(tmp_path, False, cancel)
    assert files and not results and not errors
    assert stopped and stats.cancelled


def test_xlsx_tracks_has_one_row_per_supported_file_and_status(tmp_path):
    no_tags = tmp_path / "no-tags.wav"
    corrupt = tmp_path / "corrupt.mp3"
    make_wav(no_tags)
    corrupt.write_bytes(b"bad")
    files, results, errors, _, stats = scan_library(tmp_path, False, threading.Event())
    report = export_report(tmp_path, ["Filename", "Energy", "Publisher", "Todas as tags RAW"], files, results, errors)
    book = load_workbook(report)
    tracks = list(book["Tracks"].values)
    assert len(tracks) - 1 == stats.supported == 2
    assert "Scan Status" in tracks[0]
    assert book.sheetnames == ["Tracks", "Summary", "Errors", "All Tags"]


def test_xlsx_has_no_binary_in_normalized_columns(tmp_path):
    path = tmp_path / "binary.mp3"
    save_id3(path, artwork(), TKEY(encoding=3, text=["Fm"]), comment("Energy 7"))
    result = scan_one(path)
    report = export_report(tmp_path, ["Filename", "Key", "Initial Key", "Energy", "Publisher", "Comment", "Todas as tags RAW"], [path], [(path, result)], [])
    book = load_workbook(report)
    assert sum("[BINARY DATA" in str(cell or "") for row in book["Tracks"].iter_rows(values_only=True) for cell in row) == 0
    assert sum("[BINARY DATA" in str(cell or "") for row in book["All Tags"].iter_rows(values_only=True) for cell in row) > 0


def test_xlsx_sanitizes_illegal_metadata_and_formulas(tmp_path):
    path = tmp_path / "unsafe.mp3"
    path.write_bytes(b"x")
    unsafe = "Acesse: www.packthebest.com.br -\x0bWhatsApp: (64) 099273-2714"
    raw = RawTag("COMM::eng", unsafe, "ID3", "", "eng", "COMM", False)
    result = ScanResult({"Filename": path.name, "Title": "=2+2", "Comment": unsafe, "Scan Status": "OK"}, [raw])
    report = export_report(tmp_path, ["Filename", "Title", "Comment", "Todas as tags RAW"], [path], [(path, result)], [])
    book = load_workbook(report)
    row = list(book["Tracks"].values)[1]
    assert row[2] == "'=2+2"
    assert "\\x0bWhatsApp" in row[3]
    assert "\\x0bWhatsApp" in list(book["All Tags"].values)[1][5]


def test_publisher_is_not_mass_copied_from_comment(tmp_path):
    results = []
    for index in range(5):
        path = tmp_path / f"track-{index}.mp3"
        save_id3(path, comment(f"Energy {index + 5}"))
        results.append(scan_one(path))
    assert sum(result.values["Publisher"] == result.values["Comment"] and bool(result.values["Comment"]) for result in results) == 0


def test_energy_conflict_requires_two_different_valid_numbers(tmp_path):
    path = tmp_path / "conflicts.mp3"
    save_id3(path, artwork(), comment("Energy 8"), comment("Recorded in 2018", "note", "por"))
    assert scan_one(path).values["Energy Conflict"] is False


def test_txxx_energy_level_has_priority_over_matching_comment(tmp_path):
    path = tmp_path / "energylevel-1.mp3"
    save_id3(path, TXXX(encoding=3, desc="EnergyLevel", text=["7"]), comment("Energy 7"))
    result = scan_one(path)
    assert result.values["Energy"] == "7"
    assert result.values["Energy Raw"] == "7"
    assert result.values["Energy Source"] == "TXXX:EnergyLevel"
    assert result.values["Energy Conflict"] is False


def test_txxx_energy_level_conflicts_with_comment(tmp_path):
    path = tmp_path / "energylevel-2.mp3"
    save_id3(path, TXXX(encoding=3, desc="ENERGY LEVEL", text=["7"]), comment("Energy 6"))
    result = scan_one(path)
    assert result.values["Energy Source"] == "TXXX:EnergyLevel"
    assert result.values["Energy Conflict"] is True


def test_txxx_energy_level_semicolon_values_conflict(tmp_path):
    path = tmp_path / "energylevel-3.mp3"
    save_id3(path, TXXX(encoding=3, desc="EnergyLevel", text=["7; 6"]))
    result = scan_one(path)
    assert result.values["Energy"] == "7"
    assert result.values["Energy Conflict"] is True


def test_txxx_energy_level_slash_values_conflict(tmp_path):
    path = tmp_path / "energylevel-4.mp3"
    save_id3(path, TXXX(encoding=3, desc="ENERGYLEVEL", text=["7/8"]))
    result = scan_one(path)
    assert result.values["Energy"] == "7"
    assert result.values["Energy Conflict"] is True


def test_mp4_itunes_energy_level_freeform_is_text(monkeypatch, tmp_path):
    import scanner
    class Audio:
        tags = {"----:com.apple.iTunes:energylevel": [MP4FreeForm(b"8")]}
        info = type("Info", (), {"length": 1.0})()
    monkeypatch.setattr(scanner, "File", lambda *args, **kwargs: Audio())
    path = tmp_path / "itunes.m4a"; path.write_bytes(b"fake")
    result = scan_one(path)
    assert result.values["Energy"] == "8"
    assert result.values["Energy Source"] == "----:com.apple.iTunes:energylevel"
    assert result.raw_tags[0].value == "8" and not result.raw_tags[0].is_binary


def test_mp4_mixed_in_key_energy_freeform_is_text(monkeypatch, tmp_path):
    import scanner
    class Audio:
        tags = {"----:com.mixedinkey.mixedinkey:energy": [MP4FreeForm(b"6")]}
        info = type("Info", (), {"length": 1.0})()
    monkeypatch.setattr(scanner, "File", lambda *args, **kwargs: Audio())
    path = tmp_path / "mixed-in-key.m4a"; path.write_bytes(b"fake")
    result = scan_one(path)
    assert result.values["Energy"] == "6"
    assert result.raw_tags[0].value == "6" and not result.raw_tags[0].is_binary


def test_apic_bytes_are_never_decoded_as_freeform(tmp_path):
    path = tmp_path / "cover.mp3"
    save_id3(path, artwork(1024))
    result = scan_one(path)
    assert result.raw_tags[0].is_binary
    assert result.raw_tags[0].value.startswith("[BINARY DATA")


def test_metadata_exception_uses_metadata_read_error(tmp_path):
    path = tmp_path / "bad.wav"
    path.write_bytes(b"not a wave")
    result = scan_one(path)
    assert result.error is not None
    assert result.values["Scan Status"] == "METADATA_READ_ERROR"


def test_ignored_file_is_counted_and_exported(tmp_path):
    audio = tmp_path / "audio.wav"; ignored = tmp_path / "notes.xyz"
    make_wav(audio); ignored.write_text("x")
    files, results, errors, _, stats = scan_library(tmp_path, False, threading.Event())
    assert stats.found == 2 and stats.supported == 1 and stats.ignored_extension == 1
    assert stats.ignored_extensions == {".xyz": 1}
    report = export_report(tmp_path, ["Filename"], files, results, errors, stats)
    book = load_workbook(report)
    assert "Ignored" in book.sheetnames
    assert list(book["Ignored"].values)[1] == ("notes.xyz", str(ignored.resolve()), ".xyz", "UNSUPPORTED_EXTENSION")
    summary = dict(list(book["Summary"].values)[1:6])
    assert summary["Files Found"] == 2
    assert summary["Supported Audio Files"] == 1
    assert summary["Unsupported/Ignored Files"] == 1


def test_xlsx_select_all_has_aligned_tracks_and_expected_sheets(tmp_path):
    from scanner import GROUPS
    tagged = tmp_path / "tagged.mp3"
    no_tags = tmp_path / "no-tags.wav"
    corrupt = tmp_path / "corrupt.mp3"
    ignored = tmp_path / "ignored.txt"
    save_id3(tagged, TIT2(encoding=3, text=["Title"]), TPE1(encoding=3, text=["Artist"]), artwork(), comment("Energy 7"))
    make_wav(no_tags)
    corrupt.write_bytes(b"bad")
    ignored.write_text("ignored")
    before = {path: (path.stat().st_size, path.stat().st_mtime_ns, hashlib.sha256(path.read_bytes()).hexdigest()) for path in (tagged, no_tags, corrupt)}
    files, results, errors, _, stats = scan_library(tmp_path, False, threading.Event())
    fields = [field for group in GROUPS.values() for field in group]
    report = export_report(tmp_path, fields, files, results, errors, stats)
    book = load_workbook(report, data_only=False)
    tracks = list(book["Tracks"].iter_rows(values_only=True))
    assert {"Tracks", "Summary", "Errors", "Ignored", "All Tags"}.issubset(book.sheetnames)
    assert len(tracks) == 4
    assert all(len(row) == len(tracks[0]) for row in tracks)
    assert not any(isinstance(value, str) and value.startswith("=") for sheet in book for row in sheet.iter_rows(values_only=True) for value in row)
    after = {path: (path.stat().st_size, path.stat().st_mtime_ns, hashlib.sha256(path.read_bytes()).hexdigest()) for path in (tagged, no_tags, corrupt)}
    assert before == after


def test_xlsx_partial_selection_and_missing_values_are_blank(tmp_path):
    path = tmp_path / "partial.mp3"
    path.write_bytes(b"x")
    result = ScanResult({"Filename": path.name, "Title": "Title", "Artist": "Artist", "Scan Status": "OK"})
    report = export_report(tmp_path, ["Filename", "Title", "Artist", "Energy", "Initial Key"], [path], [(path, result)], [])
    book = load_workbook(report)
    rows = list(book["Tracks"].iter_rows(values_only=True))
    record = dict(zip(rows[0], rows[1]))
    assert record["Energy"] is None
    assert record["Initial Key"] is None
    assert len(rows[0]) == len(rows[1])


def test_xlsx_without_raw_tags_omits_all_tags(tmp_path):
    path = tmp_path / "no-tags.wav"
    make_wav(path)
    files, results, errors, _, stats = scan_library(tmp_path, False, threading.Event())
    report = export_report(tmp_path, ["Filename"], files, results, errors, stats)
    book = load_workbook(report)
    assert "All Tags" not in book.sheetnames
    assert "Tracks" in book.sheetnames and "Summary" in book.sheetnames and "Errors" in book.sheetnames


def test_format_tolerates_an_empty_column_iterator():
    from app import App
    class Worksheet:
        freeze_panes = None
        auto_filter = type("AutoFilter", (), {"ref": None})()
        dimensions = "A1:A1"
        columns = [iter(())]
        column_dimensions = {}
    App._format(object.__new__(App), Worksheet())
