from __future__ import annotations
import os, re, subprocess, sys, threading, traceback, queue, webbrowser, hashlib
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from openpyxl import Workbook
from openpyxl.styles import Font
from scanner import GROUPS, DEFAULT_FIELDS, ScanStats, make_file_id, scan_library
from updater import VERSION, PROJECT_URL, check_update, download_update, launch_installer

ILLEGAL_XLSX_CHARS = re.compile(r"[\x00-\x08\x0B-\x0C\x0E-\x1F\ud800-\udfff\ufffe\uffff]")
RAW_TAGS_FIELD = "Todas as tags RAW"
FIELD_DEFINITIONS = tuple(
    (field, field)
    for group_fields in GROUPS.values()
    for field in group_fields
    if field != RAW_TAGS_FIELD
)
FIELD_DEFINITION_BY_NAME = dict(FIELD_DEFINITIONS)

def excel_safe(value):
    if not isinstance(value, str):
        return value
    value = ILLEGAL_XLSX_CHARS.sub(lambda match: f"\\x{ord(match.group()):02x}", value)
    if value.startswith(("=", "+", "-", "@")):
        value = "'" + value
    return value[:32767]

def append_safe(worksheet, values):
    worksheet.append([excel_safe(value) for value in values])

def report_file_id(path, result):
    # A removed/unmounted audio must not make an already completed scan fail.
    size = result.values.get("File Size")
    if size is not None:
        seed = str(path.resolve()).casefold() + "|" + str(size)
        return hashlib.sha1(seed.encode("utf-8", "surrogatepass")).hexdigest()[:16]
    try:
        return make_file_id(path)
    except OSError:
        return hashlib.sha1(str(path).casefold().encode("utf-8", "surrogatepass")).hexdigest()[:16]

def selected_column_definitions(fields):
    requested = set(fields)
    definitions = [definition for definition in FIELD_DEFINITIONS if definition[0] in requested]
    if "Scan Status" not in requested:
        definitions.append(("Scan Status", FIELD_DEFINITION_BY_NAME["Scan Status"]))
    return definitions

def write_error_log(output, stage, current_file, fields):
    log_path = Path(output) / "track_metadata_scanner_error.log"
    details = [
        f"Timestamp: {datetime.now().isoformat(timespec='seconds')}",
        f"Stage: {stage}",
        f"Current file: {current_file or ''}",
        f"Selected fields: {', '.join(fields)}",
        "Traceback:",
        traceback.format_exc(),
        "",
    ]
    try:
        with log_path.open("a", encoding="utf-8") as log:
            log.write("\n".join(details))
    except OSError:
        pass
    return log_path

class App(tk.Tk):
    def __init__(self):
        super().__init__(); self.title(f"Track Metadata Scanner {VERSION}"); self.geometry("900x850"); self.minsize(800, 700)
        self.events = queue.Queue(); self.dark = False; self.updating = False
        self.source = tk.StringVar(); self.output = tk.StringVar(); self.recursive = tk.BooleanVar(); self.status = tk.StringVar(value="Pronto")
        self.progress = tk.DoubleVar(); self.cancel = threading.Event(); self.worker = None; self.last_report = None; self.vars = {}
        self._build()
        self.apply_theme()
        super().after(50, self._drain_events)
        self.protocol("WM_DELETE_WINDOW", self.close_app)
    def after(self, ms, func=None, *args):
        if threading.current_thread() is not threading.main_thread():
            if func is not None: self.events.put((func, args))
            return None
        return super().after(ms, func, *args)
    def _drain_events(self):
        for _ in range(200):
            try: func, args = self.events.get_nowait()
            except queue.Empty: break
            try: func(*args)
            except Exception: self.report_callback_exception(*sys.exc_info())
        super().after(50, self._drain_events)
    def log(self, text):
        self.log_text.configure(state="normal")
        self.log_text.insert("end", f"[{datetime.now():%H:%M:%S}] {text}\n")
        if int(self.log_text.index('end-1c').split('.')[0]) > 1500:
            self.log_text.delete('1.0', '200.0')
        self.log_text.see("end"); self.log_text.configure(state="disabled")
    def close_app(self):
        if self.updating:
            messagebox.showinfo("Atualização", "Aguarde a conclusão da atualização."); return
        self.cancel.set(); self.destroy()
    def apply_theme(self):
        bg, fg, surface = ("#20242b", "#f2f4f8", "#303640") if self.dark else ("#f4f6f9", "#18202b", "#ffffff")
        style = ttk.Style(self); style.theme_use("clam")
        style.configure(".", background=bg, foreground=fg, fieldbackground=surface)
        style.configure("TEntry", fieldbackground=surface, foreground=fg)
        style.map("TButton", background=[("active", surface)])
        style.map("TCheckbutton", background=[("active", surface)])
        self.configure(bg=bg); self.canvas.configure(bg=bg)
        self.log_text.configure(bg=surface, fg=fg, insertbackground=fg)
        self.theme_btn.configure(text="☀" if self.dark else "☾")
    def toggle_theme(self):
        self.dark = not self.dark; self.apply_theme()
    def check_updates(self):
        if self.updating or (self.worker and self.worker.is_alive()): return
        self.updating = True; self.update_btn.configure(state="disabled"); self.start.configure(state="disabled")
        self.log("Consultando atualizações no GitHub...")
        def worker():
            try:
                release = check_update()
                self.after(0, lambda: self.offer_update(release))
            except Exception as exc:
                text = f"Falha ao buscar atualizações: {exc}"
                self.after(0, lambda: self.update_done(text))
        threading.Thread(target=worker, daemon=True).start()
    def update_done(self, text):
        self.updating = False; self.update_btn.configure(state="normal"); self.start.configure(state="normal"); self.log(text)
    def offer_update(self, release):
        if release is None:
            self.update_done("Nenhuma versão mais recente com executável disponível."); return
        if not getattr(sys, "frozen", False):
            self.update_done("Atualização automática disponível na versão EXE."); return
        if not messagebox.askyesno("Atualização disponível", f"Instalar {release['tag_name']} e reiniciar o programa?"):
            self.update_done("Atualização cancelada."); return
        self.log("Baixando e verificando o novo executável...")
        def worker():
            try:
                staged = download_update(release, Path(sys.executable))
                launch_installer(staged, Path(sys.executable))
                self.after(0, self.destroy)
            except Exception as exc:
                text = f"Atualização não instalada: {exc}"
                self.after(0, lambda: self.update_done(text))
        threading.Thread(target=worker, daemon=True).start()
    def _build(self):
        root = ttk.Frame(self, padding=16); root.pack(fill="both", expand=True)
        toolbar = ttk.Frame(root); toolbar.pack(fill="x")
        self.theme_btn = ttk.Button(toolbar, text="☾", width=4, command=self.toggle_theme); self.theme_btn.pack(side="right")
        self.update_btn = ttk.Button(toolbar, text="Buscar atualizações", command=self.check_updates); self.update_btn.pack(side="right", padx=6)
        ttk.Label(root, text="Track Metadata Scanner", font=("Segoe UI", 18, "bold")).pack(anchor="w")
        ttk.Label(root, text="Leitor de metadados de áudio • somente leitura", foreground="#555").pack(anchor="w", pady=(0, 12))
        for label, var in (("Pasta das músicas:", self.source), ("Pasta de saída:", self.output)):
            row = ttk.Frame(root); row.pack(fill="x", pady=3); ttk.Label(row, text=label, width=20).pack(side="left")
            ttk.Entry(row, textvariable=var).pack(side="left", fill="x", expand=True); ttk.Button(row, text="Selecionar", command=lambda v=var: self._choose(v)).pack(side="left", padx=(6,0))
        ttk.Checkbutton(root, text="Incluir subpastas", variable=self.recursive).pack(anchor="w", pady=(6, 8))
        selectrow = ttk.Frame(root); selectrow.pack(fill="x"); ttk.Label(selectrow, text="Metadados do relatório", font=("Segoe UI", 11, "bold")).pack(side="left")
        ttk.Button(selectrow, text="Selecionar tudo", command=lambda: self._all(True)).pack(side="right"); ttk.Button(selectrow, text="Limpar seleção", command=lambda: self._all(False)).pack(side="right", padx=5)
        metadata = ttk.Frame(root); metadata.pack(fill="both", expand=True)
        canvas = tk.Canvas(metadata, highlightthickness=0); self.canvas = canvas; scroll = ttk.Scrollbar(metadata, orient="vertical", command=canvas.yview); inner = ttk.Frame(canvas)
        inner.bind("<Configure>", lambda e: canvas.configure(scrollregion=canvas.bbox("all"))); canvas.create_window((0,0), window=inner, anchor="nw"); canvas.configure(yscrollcommand=scroll.set)
        canvas.pack(side="left", fill="both", expand=True, pady=8); scroll.pack(side="right", fill="y", pady=8)
        for group, fields in GROUPS.items():
            box = ttk.LabelFrame(inner, text=group, padding=6); box.pack(fill="x", pady=3)
            for i, field in enumerate(fields):
                v = tk.BooleanVar(value=field in DEFAULT_FIELDS); self.vars[field] = v
                ttk.Checkbutton(box, text=field, variable=v).grid(row=i//4, column=i%4, sticky="w", padx=8, pady=2)
        bottom = ttk.Frame(root); bottom.pack(fill="x"); ttk.Label(bottom, textvariable=self.status).pack(anchor="w")
        ttk.Progressbar(bottom, variable=self.progress, maximum=100).pack(fill="x", pady=5)
        buttons = ttk.Frame(bottom); buttons.pack(fill="x"); self.start = ttk.Button(buttons, text="INICIAR SCAN", command=self.start_scan); self.start.pack(side="left"); self.cancel_btn = ttk.Button(buttons, text="Cancelar", command=self.cancel.set, state="disabled"); self.cancel_btn.pack(side="left", padx=6); self.open_btn = ttk.Button(buttons, text="Abrir pasta do relatório", command=self.open_report, state="disabled"); self.open_btn.pack(side="right")
        ttk.Label(bottom, text="Log da análise").pack(anchor="w", pady=(8, 2))
        log_frame = ttk.Frame(bottom); log_frame.pack(fill="x")
        self.log_text = tk.Text(log_frame, height=7, wrap="word", state="disabled")
        log_scroll = ttk.Scrollbar(log_frame, command=self.log_text.yview); self.log_text.configure(yscrollcommand=log_scroll.set)
        log_scroll.pack(side="right", fill="y"); self.log_text.pack(fill="both", expand=True)
        credit = ttk.Label(bottom, text="Desenvolvido Por Pablo Escobar", cursor="hand2", foreground="#398be8")
        credit.pack(pady=(8, 0)); credit.bind("<Button-1>", lambda e: webbrowser.open(PROJECT_URL))
    def _choose(self, var):
        selected = filedialog.askdirectory();
        if selected: var.set(selected)
    def _all(self, value):
        for v in self.vars.values(): v.set(value)
    def start_scan(self):
        if self.updating or (self.worker and self.worker.is_alive()): return
        if not self.source.get().strip() or not self.output.get().strip():
            messagebox.showwarning("Pastas necessárias", "Selecione as pastas de músicas e saída."); return
        source, output = Path(self.source.get()), Path(self.output.get())
        if not source.is_dir() or not output.is_dir(): messagebox.showwarning("Pastas necessárias", "Selecione uma pasta de músicas e uma pasta de saída válidas."); return
        fields = [f for f, v in self.vars.items() if v.get()]
        if not fields: messagebox.showwarning("Seleção vazia", "Selecione pelo menos um metadado."); return
        self.cancel.clear(); self.start.config(state="disabled"); self.cancel_btn.config(state="normal"); self.open_btn.config(state="disabled"); self.status.set("Localizando arquivos..."); self.progress.set(0)
        self.log(f"Iniciando análise: {source}"); self.update_btn.configure(state="disabled")
        self.worker = threading.Thread(target=self._run, args=(source, output, fields, self.recursive.get()), daemon=True); self.worker.start()
    def _run(self, source, output, fields, recursive=False):
        stage = "Scanning metadata"
        current_file = None
        def progress(done, total, path, errors):
            nonlocal current_file
            current_file = path
            self.after(0, lambda: self.log(f"{done}/{total}: {path.name} • erros: {errors}"))
            self.after(0, lambda: (self.progress.set(done * 100 / total if total else 0), self.status.set(f"{done} / {total}  •  {path.name}  •  Erros: {errors}")))
        try:
            files, results, errors, cancelled, stats = scan_library(source, recursive, self.cancel, progress)
            for p, typ, detail in errors:
                self.after(0, lambda p=p, typ=typ: self.log(f"Erro de leitura: {p} ({typ})"))
            if errors:
                with (output / "track_metadata_scanner_error.log").open("a", encoding="utf-8") as log:
                    for p, typ, detail in errors:
                        log.write(f"\nTimestamp: {datetime.now().isoformat()}\nStage: Scanning metadata\nCurrent file: {p}\nSelected fields: {fields}\n{typ}: {detail}\n")
            if cancelled: self.after(0, lambda: self._finished("Scan cancelado", None)); return
            stage = "Exporting XLSX report"
            current_file = None
            self.after(0, lambda: self.status.set("Scan concluído. Montando relatório Excel..."))
            report = self._export(output, fields, files, results, errors, stats)
            self.after(0, lambda: self._finished(f"Scan concluído • Arquivos processados: {stats.processed} • Sem tags: {stats.no_tags} • Erros: {stats.read_errors}", report))
        except Exception as exc:
            log_path = write_error_log(output, stage, current_file, fields)
            message = f"{type(exc).__name__}: {exc}"
            self.after(0, lambda: self._finished("Falha ao gerar o relatório", None, f"{message}\nDetalhes: {log_path}"))
    def _export(self, output, fields, files, results, errors, stats=None):
        stamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S"); path = output / f"metadata_scan_{stamp}.xlsx"; n=2
        while path.exists(): path = output / f"metadata_scan_{stamp}_{n}.xlsx"; n += 1
        stats = stats or ScanStats(found=len(files), supported=len(files), processed=len(results), no_tags=sum(r.values.get("Scan Status") == "NO_TAGS" for _, r in results), read_errors=len(errors))
        wb=Workbook(); ws=wb.active; ws.title="Tracks"; columns=selected_column_definitions(fields)
        headers = ["File ID"] + [header for _, header in columns]
        file_ids = {p: report_file_id(p, result) for p, result in results}
        append_safe(ws, headers); [setattr(c, "font", Font(bold=True)) for c in ws[1]]
        for p, result in results:
            row = [file_ids[p]] + [result.values.get(key, "") for key, _ in columns]
            if len(row) != len(headers):
                raise ValueError(f"Tracks column mismatch: headers={len(headers)}, values={len(row)}, file={p}")
            append_safe(ws, row)
        self._format(ws); summary=wb.create_sheet("Summary"); append_safe(summary, ["Metric","Value"]); append_safe(summary, ["Files Found",stats.found]); append_safe(summary, ["Supported Audio Files",stats.supported]); append_safe(summary, ["Unsupported/Ignored Files",stats.ignored_extension]); append_safe(summary, ["No Tags",stats.no_tags]); append_safe(summary, ["Metadata Read Errors",stats.read_errors]); append_safe(summary, ["Arquivos processados",stats.processed]); append_safe(summary, ["Scan cancelado","TRUE" if stats.cancelled else "FALSE"])
        for field in ["Title","Artist","Genre","BPM","Key","Initial Key","Energy","Comment","ISRC"]: append_safe(summary, [f"Com {field}",sum(bool(r.values.get(field)) for _,r in results)])
        append_safe(summary, ["Com Artwork",sum(bool(r.values.get("Has Artwork")) for _,r in results)]); append_safe(summary, []); append_safe(summary, ["Extensão suportada","Quantidade"]); [append_safe(summary, [ext.upper().lstrip('.'), sum(p.suffix.casefold()==ext.casefold() for p,_ in results)]) for ext in sorted({p.suffix for p,_ in results})]; append_safe(summary, []); append_safe(summary, ["Extensão ignorada","Quantidade"]); [append_safe(summary, [ext.upper().lstrip('.'), count]) for ext, count in sorted(stats.ignored_extensions.items())]; self._format(summary)
        err=wb.create_sheet("Errors"); append_safe(err, ["Filename","Full Path","Error Type","Error Message"]); [append_safe(err, [p.name,str(p),typ,msg]) for p,typ,msg in errors]; self._format(err)
        if stats.ignored_files:
            ignored=wb.create_sheet("Ignored"); append_safe(ignored, ["Filename","Full Path","Extension","Reason"])
            for p in stats.ignored_files: append_safe(ignored, [p.name,str(p.resolve()),p.suffix.lower(),"UNSUPPORTED_EXTENSION"])
            self._format(ignored)
        if RAW_TAGS_FIELD in fields:
            raw=wb.create_sheet("All Tags"); append_safe(raw, ["File ID","Filename","Full Path","Tag Name","Tag Description","Tag Value","Source"])
            for p,r in results:
                for tag in r.raw_tags:
                    description = tag.description + (f" [lang={tag.language}]" if tag.language else "")
                    append_safe(raw, [file_ids[p],p.name,str(p),tag.name,description,tag.value,tag.source])
            self._format(raw)
        temp_path = path.with_name(path.stem + ".tmp.xlsx")
        try:
            self.after(0, lambda: self.status.set("Salvando arquivo Excel..."))
            wb.save(temp_path)
            os.replace(temp_path, path)
        finally:
            wb.close()
            if temp_path.exists():
                temp_path.unlink(missing_ok=True)
        return path
    def _format(self, ws):
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns:
            first_cell = next(iter(col), None)
            if first_cell is None:
                continue
            letter=first_cell.column_letter; ws.column_dimensions[letter].width=min(55,max(12,max(len(str(c.value or "")) for c in col)+2))
    def _finished(self, status, report, error=None):
        self.update_btn.configure(state="normal"); self.log(status + (f" • {report}" if report else "") + (f" • {error}" if error else ""))
        self.start.config(state="normal"); self.cancel_btn.config(state="disabled"); self.status.set(status + (f" • Relatório salvo em: {report}" if report else (f" • {error}" if error else ""))); self.last_report=report; self.open_btn.config(state="normal" if report else "disabled")
        if report: messagebox.showinfo("Track Metadata Scanner", status + f"\n\nRelatório salvo em:\n{report}")
        elif error: messagebox.showerror("Track Metadata Scanner", status + f"\n\n{error}")
    def open_report(self):
        if self.last_report: os.startfile(self.last_report.parent)

if __name__ == "__main__": App().mainloop()
